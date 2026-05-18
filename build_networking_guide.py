from pathlib import Path
from math import cos, sin, pi

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE_DIR = Path(r"C:\Users\yesha\Desktop\testingGPT")
ASSET_DIR = BASE_DIR / "networking_visual_assets"
OUTPUT_DOCX = BASE_DIR / "networking_teaching_guide_visual.docx"

PAGE_W = 1600
PAGE_H = 900
BG = "#F7FAFC"
NAVY = "#16324F"
BLUE = "#2C7BE5"
TEAL = "#13B0A5"
GREEN = "#4CAF50"
GOLD = "#E0A100"
ORANGE = "#F28C28"
RED = "#D64550"
PURPLE = "#6F42C1"
GRAY = "#64748B"
DARK = "#1F2937"
LIGHT = "#E5EEF8"


def ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def font(size: int, bold: bool = False):
    candidates = []
    if bold:
        candidates.extend(
            [
                r"C:\Windows\Fonts\arialbd.ttf",
                r"C:\Windows\Fonts\calibrib.ttf",
                r"C:\Windows\Fonts\segoeuib.ttf",
            ]
        )
    candidates.extend(
        [
            r"C:\Windows\Fonts\arial.ttf",
            r"C:\Windows\Fonts\calibri.ttf",
            r"C:\Windows\Fonts\segoeui.ttf",
        ]
    )
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


TITLE_FONT = font(48, bold=True)
SUBTITLE_FONT = font(30, bold=True)
BODY_FONT = font(24)
SMALL_FONT = font(20)
TINY_FONT = font(18)


def new_canvas():
    img = Image.new("RGB", (PAGE_W, PAGE_H), BG)
    draw = ImageDraw.Draw(img)
    return img, draw


def rounded_box(draw, xy, fill, outline=None, width=3, radius=24):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def centered_text(draw, xy, text, font_obj, fill=DARK):
    bbox = draw.textbbox((0, 0), text, font=font_obj)
    x = xy[0] - (bbox[2] - bbox[0]) / 2
    y = xy[1] - (bbox[3] - bbox[1]) / 2
    draw.text((x, y), text, font=font_obj, fill=fill)


def multiline(draw, x, y, text, font_obj, fill=DARK, spacing=10):
    draw.multiline_text((x, y), text, font=font_obj, fill=fill, spacing=spacing)


def arrow(draw, start, end, fill=NAVY, width=8, head=18):
    draw.line([start, end], fill=fill, width=width)
    angle = pi + atan2(end[1] - start[1], end[0] - start[0])
    left = (end[0] + head * cos(angle - pi / 8), end[1] + head * sin(angle - pi / 8))
    right = (end[0] + head * cos(angle + pi / 8), end[1] + head * sin(angle + pi / 8))
    draw.polygon([end, left, right], fill=fill)


def atan2(y, x):
    from math import atan2 as _atan2

    return _atan2(y, x)


def device(draw, x, y, w, h, label, fill=LIGHT, outline=BLUE):
    rounded_box(draw, (x, y, x + w, y + h), fill=fill, outline=outline, width=4, radius=18)
    centered_text(draw, (x + w / 2, y + h / 2), label, BODY_FONT)


def label_chip(draw, x, y, text, fill=BLUE):
    bbox = draw.textbbox((0, 0), text, font=SMALL_FONT)
    pad_x = 16
    pad_y = 10
    rounded_box(draw, (x, y, x + bbox[2] + pad_x * 2, y + bbox[3] + pad_y * 2), fill=fill, radius=18)
    draw.text((x + pad_x, y + pad_y - 2), text, font=SMALL_FONT, fill="white")


def title_block(draw, title, subtitle):
    draw.rectangle((0, 0, PAGE_W, 110), fill=NAVY)
    draw.text((55, 28), title, font=TITLE_FONT, fill="white")
    draw.text((58, 122), subtitle, font=SMALL_FONT, fill=GRAY)


def draw_stack(draw, x, y, width, height, labels, colors):
    layer_h = height / len(labels)
    for idx, (label, color) in enumerate(zip(labels, colors)):
        top = y + idx * layer_h
        rounded_box(draw, (x, top, x + width, top + layer_h - 10), fill=color, outline="white", width=2, radius=18)
        centered_text(draw, (x + width / 2, top + layer_h / 2 - 5), label, SUBTITLE_FONT, fill="white")


def save_diagram(name, painter):
    img, draw = new_canvas()
    painter(img, draw)
    path = ASSET_DIR / f"{name}.png"
    img.save(path)
    return path


def make_diagrams():
    ensure_dir(ASSET_DIR)
    diagrams = []

    def d01(img, draw):
        title_block(draw, "1. What a Network Does", "Devices exchange data, services, and internet access through shared links.")
        items = [
            (170, 260, "Laptop"),
            (470, 260, "Phone"),
            (770, 260, "Printer"),
            (1070, 260, "Server"),
            (470, 560, "Router"),
            (870, 560, "Internet"),
        ]
        for x, y, label in items:
            device(draw, x, y, 190, 110, label, fill="#FFFFFF")
        for start, end in [
            ((265, 370), (565, 560)),
            ((565, 370), (565, 560)),
            ((865, 370), (665, 560)),
            ((1165, 370), (665, 560)),
            ((665, 615), (870, 615)),
        ]:
            arrow(draw, start, end, fill=TEAL)
        multiline(draw, 1000, 520, "Shared uses:\n- Files\n- Web access\n- Printing\n- Applications", BODY_FONT)

    diagrams.append(("what_network_does", "A network lets devices share resources and services.", save_diagram("01_what_network_does", d01)))

    def d02(img, draw):
        title_block(draw, "2. PAN, LAN, MAN, WAN", "Networks grow from personal range to global reach.")
        boxes = [
            (120, 250, 250, 180, GREEN, "PAN", "Personal devices\nwithin a few meters"),
            (430, 220, 300, 240, BLUE, "LAN", "Home, school,\nor office network"),
            (810, 190, 320, 300, PURPLE, "MAN", "City-wide\nnetwork coverage"),
            (1180, 150, 280, 380, ORANGE, "WAN", "Countries,\ncontinents,\ninternet"),
        ]
        for x, y, w, h, color, name, note in boxes:
            rounded_box(draw, (x, y, x + w, y + h), fill=color, radius=28)
            centered_text(draw, (x + w / 2, y + 55), name, TITLE_FONT, fill="white")
            centered_text(draw, (x + w / 2, y + h / 2 + 30), note, BODY_FONT, fill="white")
        arrow(draw, (370, 340), (430, 340), fill=NAVY)
        arrow(draw, (730, 340), (810, 340), fill=NAVY)
        arrow(draw, (1130, 340), (1180, 340), fill=NAVY)

    diagrams.append(("network_types", "PAN, LAN, MAN, and WAN describe network size and reach.", save_diagram("02_network_types", d02)))

    def d03(img, draw):
        title_block(draw, "3. Physical Topologies", "The arrangement of links changes cost, reliability, and troubleshooting.")
        # Star
        centered_text(draw, (250, 180), "Star", SUBTITLE_FONT)
        center = (250, 380)
        device(draw, 190, 340, 120, 80, "Switch", fill="#FFFFFF")
        pts = [(250, 250), (120, 380), (250, 520), (380, 380)]
        for px, py in pts:
            device(draw, px - 45, py - 35, 90, 70, "PC", fill="#FFFFFF")
            arrow(draw, center, (px, py), fill=BLUE)
        # Bus
        centered_text(draw, (800, 180), "Bus", SUBTITLE_FONT)
        draw.line((600, 390, 1000, 390), fill=TEAL, width=10)
        for px in [640, 760, 880, 960]:
            draw.line((px, 300, px, 390), fill=GRAY, width=6)
            device(draw, px - 45, 220, 90, 70, "PC", fill="#FFFFFF")
        # Ring
        centered_text(draw, (1310, 180), "Ring", SUBTITLE_FONT)
        points = []
        cx, cy, r = 1310, 390, 140
        for i in range(6):
            a = i * 2 * pi / 6 - pi / 2
            points.append((cx + r * cos(a), cy + r * sin(a)))
        draw.line(points + [points[0]], fill=ORANGE, width=10)
        for px, py in points:
            device(draw, px - 40, py - 30, 80, 60, "PC", fill="#FFFFFF")
        multiline(draw, 140, 640, "Star: easy to manage\nBus: simple but fragile\nRing: predictable flow", BODY_FONT)

    diagrams.append(("topologies_basic", "Common topologies change how devices are arranged and how failures spread.", save_diagram("03_topologies_basic", d03)))

    def d04(img, draw):
        title_block(draw, "4. Mesh and Tree Topologies", "Large networks often need multiple paths or hierarchical structure.")
        centered_text(draw, (430, 180), "Mesh", SUBTITLE_FONT)
        pts = [(250, 320), (430, 250), (620, 320), (300, 520), (560, 520)]
        for i, a in enumerate(pts):
            for b in pts[i + 1 :]:
                draw.line([a, b], fill=BLUE, width=4)
        for px, py in pts:
            device(draw, px - 42, py - 30, 84, 60, "Node", fill="#FFFFFF")
        centered_text(draw, (1160, 180), "Tree", SUBTITLE_FONT)
        root = (1160, 240)
        mids = [(980, 390), (1160, 390), (1340, 390)]
        leaves = [(900, 560), (1030, 560), (1110, 560), (1210, 560), (1290, 560), (1420, 560)]
        device(draw, root[0] - 60, root[1] - 35, 120, 70, "Core", fill="#FFFFFF")
        for m in mids:
            arrow(draw, root, m, fill=GREEN)
            device(draw, m[0] - 55, m[1] - 35, 110, 70, "Switch", fill="#FFFFFF")
        for idx, leaf in enumerate(leaves):
            parent = mids[idx // 2]
            arrow(draw, parent, leaf, fill=GREEN)
            device(draw, leaf[0] - 45, leaf[1] - 30, 90, 60, "Host", fill="#FFFFFF")
        multiline(draw, 180, 670, "Mesh offers resilience.\nTree offers clean scaling for schools and offices.", BODY_FONT)

    diagrams.append(("topologies_advanced", "Mesh favors resilience; tree favors organized growth.", save_diagram("04_topologies_advanced", d04)))

    def d05(img, draw):
        title_block(draw, "5. Core Network Devices", "Each device plays a different role in moving or protecting traffic.")
        cards = [
            (120, 220, BLUE, "NIC", "Gives a device\nnetwork access"),
            (390, 220, GREEN, "Switch", "Moves frames\ninside a LAN"),
            (660, 220, ORANGE, "Router", "Connects\ndifferent networks"),
            (930, 220, PURPLE, "Access Point", "Adds wireless\nconnectivity"),
            (1200, 220, RED, "Firewall", "Filters and\nblocks traffic"),
            (390, 520, TEAL, "Modem", "Links your site\nto the ISP"),
            (760, 520, GOLD, "Server", "Provides apps,\nfiles, or websites"),
        ]
        for x, y, color, title, note in cards:
            rounded_box(draw, (x, y, x + 220, y + 170), fill=color, radius=24)
            centered_text(draw, (x + 110, y + 45), title, SUBTITLE_FONT, fill="white")
            centered_text(draw, (x + 110, y + 108), note, BODY_FONT, fill="white")

    diagrams.append(("core_devices", "Switches, routers, firewalls, and access points each solve different jobs.", save_diagram("05_core_devices", d05)))

    def d06(img, draw):
        title_block(draw, "6. Transmission Media", "Media choices affect speed, cost, reliability, and mobility.")
        media = [
            (140, 240, 360, 420, BLUE, "Twisted Pair", "Common Ethernet cable\nLow cost\nGood for offices"),
            (620, 240, 360, 420, GREEN, "Fiber Optic", "Light-based transmission\nVery high speed\nLong distance"),
            (1100, 240, 360, 420, ORANGE, "Wireless", "Wi-Fi, Bluetooth,\ncellular links\nFlexible movement"),
        ]
        for x, y, w, h, color, name, note in media:
            rounded_box(draw, (x, y, x + w, y + h), fill="#FFFFFF", outline=color, width=6, radius=26)
            draw.rectangle((x + 24, y + 24, x + w - 24, y + 110), fill=color)
            centered_text(draw, (x + w / 2, y + 68), name, SUBTITLE_FONT, fill="white")
            centered_text(draw, (x + w / 2, y + 270), note, BODY_FONT)

    diagrams.append(("media", "Wired and wireless media carry traffic with different tradeoffs.", save_diagram("06_media", d06)))

    def d07(img, draw):
        title_block(draw, "7. Client-Server Model", "Most modern services work as clients requesting resources from servers.")
        device(draw, 150, 310, 220, 140, "Client 1\nBrowser", fill="#FFFFFF")
        device(draw, 150, 520, 220, 140, "Client 2\nApp", fill="#FFFFFF")
        device(draw, 1160, 400, 260, 180, "Server\nWeb / File / Mail", fill="#FFFFFF", outline=PURPLE)
        device(draw, 620, 400, 220, 140, "Switch / LAN", fill="#FFFFFF", outline=GREEN)
        for start in [(370, 380), (370, 590)]:
            arrow(draw, start, (620, 470), fill=BLUE)
            arrow(draw, (840, 470), (1160, 490), fill=PURPLE)
        arrow(draw, (1160, 540), (840, 520), fill=PURPLE)
        arrow(draw, (620, 520), (370, 640), fill=BLUE)
        multiline(draw, 470, 200, "Clients ask for services.\nServers respond with data or applications.", SUBTITLE_FONT)

    diagrams.append(("client_server", "Clients send requests; servers answer with services and data.", save_diagram("07_client_server", d07)))

    def d08(img, draw):
        title_block(draw, "8. Peer-to-Peer Model", "In peer-to-peer networking, each device can both request and provide resources.")
        peers = [(350, 250), (900, 250), (1250, 540), (550, 620), (180, 500)]
        for idx, (px, py) in enumerate(peers, start=1):
            device(draw, px, py, 160, 90, f"Peer {idx}", fill="#FFFFFF", outline=TEAL)
        for i, a in enumerate(peers):
            for b in peers[i + 1 :]:
                arrow(draw, (a[0] + 80, a[1] + 45), (b[0] + 80, b[1] + 45), fill=GRAY, width=4, head=12)
        multiline(draw, 520, 120, "Peers share directly.\nGood for small groups, less ideal for central control.", SUBTITLE_FONT)

    diagrams.append(("peer_to_peer", "Peer-to-peer networks share resources directly without a central server.", save_diagram("08_peer_to_peer", d08)))

    def d09(img, draw):
        title_block(draw, "9. Home Network Layout", "A simple home or classroom network shows the full path from ISP to user devices.")
        device(draw, 120, 360, 180, 100, "ISP", fill="#FFFFFF", outline=ORANGE)
        device(draw, 370, 360, 180, 100, "Modem", fill="#FFFFFF", outline=TEAL)
        device(draw, 620, 330, 220, 150, "Wi-Fi Router\nDHCP + NAT + Firewall", fill="#FFFFFF", outline=BLUE)
        device(draw, 980, 200, 180, 90, "Laptop", fill="#FFFFFF")
        device(draw, 1240, 200, 180, 90, "Phone", fill="#FFFFFF")
        device(draw, 980, 500, 180, 90, "Smart TV", fill="#FFFFFF")
        device(draw, 1240, 500, 180, 90, "Printer", fill="#FFFFFF")
        arrow(draw, (300, 410), (370, 410), fill=ORANGE)
        arrow(draw, (550, 410), (620, 410), fill=TEAL)
        for target in [(980, 245), (1240, 245), (980, 545), (1240, 545)]:
            arrow(draw, (840, 405), target, fill=BLUE)
        label_chip(draw, 670, 520, "Default Gateway", fill=PURPLE)

    diagrams.append(("home_network", "A home network typically centers around one router providing multiple services.", save_diagram("09_home_network", d09)))

    def d10(img, draw):
        title_block(draw, "10. OSI Model", "The OSI model breaks communication into seven learning layers.")
        draw_stack(
            draw,
            500,
            180,
            600,
            620,
            [
                "7 Application",
                "6 Presentation",
                "5 Session",
                "4 Transport",
                "3 Network",
                "2 Data Link",
                "1 Physical",
            ],
            [PURPLE, BLUE, TEAL, GREEN, GOLD, ORANGE, RED],
        )
        multiline(draw, 120, 270, "Top layers:\nuser-facing services\n\nMiddle:\nreliability and routing\n\nBottom:\nlocal delivery and signals", SUBTITLE_FONT)

    diagrams.append(("osi_model", "OSI is a teaching model that explains communication layer by layer.", save_diagram("10_osi_model", d10)))

    def d11(img, draw):
        title_block(draw, "11. TCP/IP Model", "The internet mostly uses the simpler TCP/IP model.")
        draw_stack(
            draw,
            540,
            220,
            520,
            500,
            [
                "Application",
                "Transport",
                "Internet",
                "Network Access",
            ],
            [PURPLE, GREEN, ORANGE, BLUE],
        )
        multiline(draw, 180, 280, "Application: HTTP, DNS, SMTP\nTransport: TCP, UDP\nInternet: IP, ICMP\nNetwork Access: Ethernet, Wi-Fi", SUBTITLE_FONT)

    diagrams.append(("tcp_ip_model", "TCP/IP is the practical model used by the internet.", save_diagram("11_tcp_ip_model", d11)))

    def d12(img, draw):
        title_block(draw, "12. Encapsulation", "Each layer adds its own information before data travels across the network.")
        layers = [
            ("Application Data", PURPLE),
            ("TCP Segment", GREEN),
            ("IP Packet", ORANGE),
            ("Ethernet Frame", BLUE),
            ("Bits on Media", RED),
        ]
        x, y, w, h = 260, 220, 1080, 110
        for idx, (label, color) in enumerate(layers):
            pad = idx * 70
            rounded_box(draw, (x + pad, y + idx * 85, x + w - pad, y + h + idx * 85), fill=color, radius=24)
            centered_text(draw, ((x + w) / 2, y + h / 2 + idx * 85), label, SUBTITLE_FONT, fill="white")

    diagrams.append(("encapsulation", "Data becomes segments, packets, frames, and finally signals.", save_diagram("12_encapsulation", d12)))

    def d13(img, draw):
        title_block(draw, "13. IPv4 Addressing", "An IPv4 address uses four decimal octets to identify a host.")
        rounded_box(draw, (220, 290, 1380, 470), fill="#FFFFFF", outline=BLUE, width=5, radius=30)
        octets = ["192", "168", "1", "10"]
        start_x = 280
        for idx, octet in enumerate(octets):
            x = start_x + idx * 250
            rounded_box(draw, (x, 330, x + 160, 420), fill=LIGHT, outline=TEAL, width=4, radius=20)
            centered_text(draw, (x + 80, 375), octet, TITLE_FONT)
            if idx < 3:
                centered_text(draw, (x + 190, 374), ".", TITLE_FONT)
        label_chip(draw, 340, 520, "Network portion", fill=GREEN)
        label_chip(draw, 980, 520, "Host portion", fill=ORANGE)
        draw.line((360, 510, 830, 510), fill=GREEN, width=8)
        draw.line((1000, 510, 1160, 510), fill=ORANGE, width=8)
        multiline(draw, 420, 650, "Private ranges include 192.168.x.x, 10.x.x.x, and 172.16-31.x.x.", SUBTITLE_FONT)

    diagrams.append(("ipv4", "IPv4 addresses use four octets with network and host portions.", save_diagram("13_ipv4", d13)))

    def d14(img, draw):
        title_block(draw, "14. Subnet Mask and Gateway", "The subnet mask shows the local network; the gateway reaches everything else.")
        device(draw, 170, 360, 260, 130, "PC\n192.168.1.10/24", fill="#FFFFFF")
        device(draw, 620, 360, 320, 130, "Router / Gateway\n192.168.1.1", fill="#FFFFFF", outline=BLUE)
        device(draw, 1120, 360, 260, 130, "Internet Site\n8.8.8.8", fill="#FFFFFF", outline=ORANGE)
        arrow(draw, (430, 425), (620, 425), fill=GREEN)
        arrow(draw, (940, 425), (1120, 425), fill=ORANGE)
        multiline(draw, 220, 590, "/24 means 255.255.255.0\n192.168.1 is local.\nAnything outside goes to the default gateway.", SUBTITLE_FONT)

    diagrams.append(("subnet_gateway", "Subnet masks define the local range; gateways forward to remote networks.", save_diagram("14_subnet_gateway", d14)))

    def d15(img, draw):
        title_block(draw, "15. DNS Name Resolution", "DNS translates readable names into IP addresses.")
        chain = [
            ("User types\nwww.example.com", 120, BLUE),
            ("Resolver\nasks DNS", 450, GREEN),
            ("Authoritative DNS\nreturns IP", 810, PURPLE),
            ("Browser reaches\nserver", 1180, ORANGE),
        ]
        for label, x, color in chain:
            device(draw, x, 330, 230, 140, label, fill="#FFFFFF", outline=color)
        for x in [350, 710, 1070]:
            arrow(draw, (x, 400), (x + 100, 400), fill=NAVY)
        multiline(draw, 430, 580, "Example result: www.example.com -> 93.184.216.34", SUBTITLE_FONT)

    diagrams.append(("dns", "DNS is the internet's naming system.", save_diagram("15_dns", d15)))

    def d16(img, draw):
        title_block(draw, "16. DHCP Address Assignment", "DHCP gives devices network settings automatically through the DORA sequence.")
        device(draw, 180, 340, 220, 140, "Client", fill="#FFFFFF", outline=BLUE)
        device(draw, 1180, 340, 220, 140, "DHCP Server", fill="#FFFFFF", outline=PURPLE)
        steps = [("Discover", 270), ("Offer", 520), ("Request", 770), ("Acknowledge", 1020)]
        for idx, (label, x) in enumerate(steps):
            y = 280 if idx % 2 == 0 else 520
            arrow(draw, (400 if idx % 2 == 0 else 1180, 410), (x, y + 40), fill=GREEN if idx % 2 == 0 else ORANGE)
            label_chip(draw, x - 40, y, label, fill=GREEN if idx % 2 == 0 else ORANGE)
        multiline(draw, 460, 660, "DHCP usually provides IP address, subnet mask, gateway, and DNS server.", SUBTITLE_FONT)

    diagrams.append(("dhcp", "DHCP automates IP setup through Discover, Offer, Request, and Acknowledge.", save_diagram("16_dhcp", d16)))

    def d17(img, draw):
        title_block(draw, "17. Switch Forwarding", "A switch learns MAC addresses and sends frames only to the right port.")
        device(draw, 610, 310, 280, 180, "Switch\nMAC Table:\nA1 -> Port 1\nB2 -> Port 3", fill="#FFFFFF", outline=GREEN)
        hosts = [
            (150, 220, "Host A\nMAC A1"),
            (150, 540, "Host B\nMAC B2"),
            (1170, 220, "Host C\nMAC C3"),
            (1170, 540, "Host D\nMAC D4"),
        ]
        anchors = [(430, 310), (430, 490), (1170, 310), (1170, 490)]
        for (x, y, label), (ax, ay) in zip(hosts, anchors):
            device(draw, x, y, 220, 110, label, fill="#FFFFFF")
            arrow(draw, (x + 220, y + 55) if x < 600 else (x, y + 55), (610 if x < 600 else 890, ay), fill=BLUE)
        label_chip(draw, 950, 650, "Forward only to the destination port", fill=TEAL)

    diagrams.append(("switch", "Switches use MAC tables to avoid sending every frame everywhere.", save_diagram("17_switch", d17)))

    def d18(img, draw):
        title_block(draw, "18. Routing Between Networks", "Routers connect separate IP networks and choose the next hop.")
        nets = [
            (120, 330, 260, 180, BLUE, "LAN 1\n192.168.1.0/24"),
            (660, 280, 280, 240, GREEN, "Router\nRouting table decides\nwhere traffic goes"),
            (1180, 180, 250, 150, ORANGE, "LAN 2\n10.0.0.0/24"),
            (1180, 470, 250, 150, PURPLE, "Internet\nremote networks"),
        ]
        for x, y, w, h, color, label in nets:
            rounded_box(draw, (x, y, x + w, y + h), fill="#FFFFFF", outline=color, width=5, radius=24)
            centered_text(draw, (x + w / 2, y + h / 2), label, BODY_FONT)
        arrow(draw, (380, 420), (660, 400), fill=BLUE)
        arrow(draw, (940, 360), (1180, 255), fill=ORANGE)
        arrow(draw, (940, 460), (1180, 545), fill=PURPLE)

    diagrams.append(("routing", "Routers move packets between different IP networks.", save_diagram("18_routing", d18)))

    def d19(img, draw):
        title_block(draw, "19. TCP and UDP", "Transport protocols trade reliability against speed and simplicity.")
        left = (180, 220, 680, 720)
        right = (920, 220, 1420, 720)
        rounded_box(draw, left, fill="#FFFFFF", outline=GREEN, width=6, radius=28)
        rounded_box(draw, right, fill="#FFFFFF", outline=ORANGE, width=6, radius=28)
        centered_text(draw, (430, 290), "TCP", TITLE_FONT, fill=GREEN)
        centered_text(draw, (1170, 290), "UDP", TITLE_FONT, fill=ORANGE)
        multiline(draw, 240, 360, "- Connection-oriented\n- Reliable and ordered\n- Retransmits lost data\n- Good for web and email", BODY_FONT)
        multiline(draw, 980, 360, "- Connectionless\n- Lower overhead\n- Faster delivery\n- Good for voice, video, gaming", BODY_FONT)

    diagrams.append(("tcp_udp", "TCP emphasizes reliability; UDP emphasizes speed.", save_diagram("19_tcp_udp", d19)))

    def d20(img, draw):
        title_block(draw, "20. Ports and Services", "Ports tell a device which application should receive the traffic.")
        device(draw, 180, 280, 240, 340, "Server\nIP 192.168.1.20", fill="#FFFFFF", outline=BLUE)
        ports = [("80", "HTTP"), ("443", "HTTPS"), ("25", "SMTP"), ("53", "DNS"), ("22", "SSH")]
        y = 250
        for num, name in ports:
            rounded_box(draw, (560, y, 1180, y + 80), fill="#FFFFFF", outline=TEAL, width=4, radius=18)
            centered_text(draw, (650, y + 40), num, SUBTITLE_FONT, fill=BLUE)
            draw.text((760, y + 18), name, font=SUBTITLE_FONT, fill=DARK)
            arrow(draw, (420, 450), (560, y + 40), fill=GRAY, width=4, head=12)
            y += 95

    diagrams.append(("ports", "Ports separate web, mail, DNS, and other services on the same host.", save_diagram("20_ports", d20)))

    def d21(img, draw):
        title_block(draw, "21. Packet Journey to a Website", "A browser request touches naming, routing, transport, and response delivery.")
        stages = [
            ("1 Browser\nrequest", 120, BLUE),
            ("2 DNS\nlookup", 400, GREEN),
            ("3 Router\nforwards", 680, ORANGE),
            ("4 Web server\nresponds", 960, PURPLE),
            ("5 Page shown\nto user", 1240, TEAL),
        ]
        for label, x, color in stages:
            device(draw, x, 330, 180, 150, label, fill="#FFFFFF", outline=color)
        for x in [300, 580, 860, 1140]:
            arrow(draw, (x, 405), (x + 100, 405), fill=NAVY)

    diagrams.append(("packet_journey", "Opening one site involves several coordinated networking steps.", save_diagram("21_packet_journey", d21)))

    def d22(img, draw):
        title_block(draw, "22. Security Basics", "Security layers protect users, devices, and data as traffic moves.")
        rings = [
            ((800, 430), 300, RED, "Firewall"),
            ((800, 430), 210, ORANGE, "Encryption"),
            ((800, 430), 120, GREEN, "Authentication"),
        ]
        for center, radius, color, label in rings:
            draw.ellipse((center[0] - radius, center[1] - radius, center[0] + radius, center[1] + radius), outline=color, width=12)
            centered_text(draw, (center[0], center[1] - radius + 40), label, SUBTITLE_FONT, fill=color)
        centered_text(draw, (800, 430), "User Data", TITLE_FONT)
        multiline(draw, 180, 690, "Add strong passwords, WPA2/WPA3, patching, antivirus, and phishing awareness on top.", SUBTITLE_FONT)

    diagrams.append(("security", "Networking security combines filtering, encryption, and identity checks.", save_diagram("22_security", d22)))

    return diagrams


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def build_doc(diagrams):
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(24)
    styles["Title"].font.bold = True
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(18)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].font.color.rgb = RGBColor(22, 50, 79)
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(14)
    styles["Heading 2"].font.bold = True
    styles["Heading 2"].font.color.rgb = RGBColor(44, 123, 229)
    if "Caption" not in styles:
        styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    styles["Caption"].font.name = "Arial"
    styles["Caption"].font.size = Pt(9)
    styles["Caption"].font.italic = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Networking Basics Teaching Guide")
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(22, 50, 79)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("A detailed beginner-friendly handbook with visual explanations for classroom teaching")
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(100, 116, 139)

    info = doc.add_table(rows=2, cols=2)
    info.style = "Table Grid"
    info.cell(0, 0).text = "Audience"
    info.cell(0, 1).text = "Students and beginners learning core networking concepts"
    info.cell(1, 0).text = "Coverage"
    info.cell(1, 1).text = "Definitions, models, addressing, protocols, traffic flow, security, and troubleshooting"
    for row in info.rows:
        row.cells[0].width = Inches(1.5)
        row.cells[1].width = Inches(4.8)
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for row in [info.rows[0], info.rows[1]]:
        set_cell_shading(row.cells[0], "DCE9F8")

    doc.add_paragraph("")
    doc.add_heading("How to Use This Guide", level=1)
    for line in [
        "Teach from simple to complex: start with what a network is, then move into addressing, protocols, and troubleshooting.",
        "Use the diagrams to pause and explain the flow, not just the definitions. Students usually remember pictures faster than terms.",
        "The guide is written so you can teach it as a single session or split it across multiple classes.",
    ]:
        para = doc.add_paragraph(style="List Bullet")
        para.add_run(line)

    doc.add_heading("Suggested Teaching Sequence", level=1)
    for idx, line in enumerate(
        [
            "Introduction to networking and why it matters",
            "Types of networks and topologies",
            "Devices, media, and basic communication flow",
            "IP addressing, subnetting basics, DNS, and DHCP",
            "OSI model, TCP/IP model, and transport protocols",
            "Ports, packet delivery, security, and troubleshooting",
        ],
        start=1,
    ):
        doc.add_paragraph(f"{idx}. {line}")

    sections = [
        (
            "Networking Fundamentals",
            [
                ("What Networking Means", "Networking is the process of connecting devices so they can exchange data and share services such as internet access, files, printers, and applications. The idea is simple: one device creates or requests data, and another device receives or provides it."),
                ("Where Students See It", "Students already use networking every day when they browse a website, join a video call, print from a lab computer, stream content, or send a message from a phone. Anchoring the topic in these normal actions makes the technical terms easier to understand."),
            ],
            diagrams[0:2],
        ),
        (
            "Layouts and Devices",
            [
                ("Topologies", "A topology is the shape or arrangement of a network. It affects cost, maintenance, growth, and what happens when one link fails. Star topology is common in modern LANs because it is easier to manage than bus or ring layouts."),
                ("Devices", "Devices in a network have different jobs. A switch moves traffic inside the same LAN, a router connects one network to another, a modem links the local site to the internet provider, and an access point adds Wi-Fi. A firewall applies security rules to traffic."),
                ("Media", "Traffic can travel over copper cable, fiber optic cable, or wireless radio waves. Wired media is usually steadier and faster. Wireless media is more flexible and easier to move around with."),
            ],
            diagrams[2:6],
        ),
        (
            "Models and Communication Patterns",
            [
                ("Client-Server and Peer-to-Peer", "In the client-server model, users ask for resources from a central server. In peer-to-peer networking, devices can share directly with each other. Client-server is more common in schools, offices, and web applications because it is easier to secure and manage."),
                ("OSI and TCP/IP", "The OSI model is a teaching model with seven layers, while the TCP/IP model is the practical model used by the internet. Both models help students understand where each protocol belongs and where problems may occur."),
                ("Encapsulation", "As data moves down the stack, each layer adds its own control information. This is why the same message is called data, then a segment, then a packet, then a frame depending on the layer being discussed."),
            ],
            diagrams[6:12],
        ),
        (
            "Addressing and Core Services",
            [
                ("IP and MAC Addresses", "An IP address is a logical address used to identify a host on a network, while a MAC address is the physical hardware address used mainly for local delivery. A helpful teaching line is: IP says where the device is, MAC says who the device is."),
                ("Subnet Mask and Gateway", "The subnet mask separates the network portion from the host portion of an address. The default gateway is the route a device uses when the destination is outside the local network."),
                ("DNS and DHCP", "DNS converts human-friendly names into IP addresses, and DHCP automatically provides IP settings such as address, mask, gateway, and DNS server. These two services remove a huge amount of manual effort from networking."),
            ],
            diagrams[12:16],
        ),
        (
            "Traffic Movement and Protocols",
            [
                ("Switching and Routing", "Switching happens inside a local network and relies heavily on MAC addresses. Routing happens between different networks and relies on IP addresses and routing decisions."),
                ("TCP and UDP", "TCP builds a reliable conversation with acknowledgements, retransmissions, and ordered delivery. UDP sends quickly with less overhead, which is why it is useful for live audio, video, and gaming."),
                ("Ports and Sockets", "Ports let a single device offer many services at once. A socket combines an IP address and port number, such as 192.168.1.20:443."),
                ("Packet Journey", "When a user opens a website, the browser may query DNS, the host creates packets, routers forward the traffic, the server responds, and the browser rebuilds the content into a visible page."),
            ],
            diagrams[16:21],
        ),
        (
            "Security and Troubleshooting",
            [
                ("Security Basics", "Firewalls, encryption, authentication, strong passwords, safe Wi-Fi settings, and phishing awareness all support secure networking. Students should understand that networking is not only about connection, but also about controlled and trusted connection."),
                ("Troubleshooting Flow", "Basic troubleshooting starts with physical checks, then IP settings, then reachability tests such as ping, then DNS checks, and finally service or firewall review. Teaching a repeatable method is often more valuable than teaching isolated commands."),
            ],
            diagrams[21:22],
        ),
    ]

    diagram_counter = 1
    for section_title, paragraphs, diagram_slice in sections:
        doc.add_page_break()
        doc.add_heading(section_title, level=1)
        for subheading, text in paragraphs:
            doc.add_heading(subheading, level=2)
            doc.add_paragraph(text)
        for key, caption, path in diagram_slice:
            doc.add_picture(str(path), width=Inches(6.6))
            cap = doc.add_paragraph(style="Caption")
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.add_run(f"Figure {diagram_counter}. {caption}")
            diagram_counter += 1
            doc.add_paragraph("")

    doc.add_page_break()
    doc.add_heading("Quick Revision Table", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Term"
    hdr[1].text = "Short Meaning"
    hdr[2].text = "Example"
    set_repeat_table_header(table.rows[0])
    for cell in hdr:
        set_cell_shading(cell, "16324F")
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.bold = True
    revisions = [
        ("LAN", "A local network inside a limited area", "School lab"),
        ("WAN", "A large network across long distances", "Internet"),
        ("Switch", "Connects devices within a LAN", "Computer to printer in same office"),
        ("Router", "Connects different networks", "Home router to ISP"),
        ("DNS", "Turns names into IP addresses", "google.com -> IP"),
        ("DHCP", "Assigns network settings automatically", "Phone joins Wi-Fi"),
        ("TCP", "Reliable transport protocol", "Web page loading"),
        ("UDP", "Fast transport protocol", "Online game traffic"),
        ("Firewall", "Filters network traffic", "Blocks unsafe access"),
    ]
    for term, meaning, example in revisions:
        row = table.add_row().cells
        row[0].text = term
        row[1].text = meaning
        row[2].text = example
        for cell in row:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    doc.add_heading("Classroom Closing Summary", level=1)
    doc.add_paragraph(
        "A strong beginner foundation in networking should answer five big questions: what devices are connected, how they are arranged, how they are addressed, how data moves, and how the network is protected. Once students are confident with those five, advanced topics such as subnetting, VLANs, routing protocols, and cloud networking become much easier to learn."
    )

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_p.add_run("Networking Basics Teaching Guide")
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor(100, 116, 139)

    doc.save(OUTPUT_DOCX)


def main():
    diagrams = make_diagrams()
    build_doc(diagrams)
    print(f"Created {OUTPUT_DOCX}")
    print(f"Created {len(diagrams)} diagrams in {ASSET_DIR}")


if __name__ == "__main__":
    main()
